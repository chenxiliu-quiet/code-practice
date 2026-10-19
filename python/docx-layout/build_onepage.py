# -*- coding: utf-8 -*-
"""把作业 docx 重排成「一页一题」：每题从新页开始，图片等比缩放以适配单页。"""
import re
import sys
import copy
from docx import Document
from docx.shared import Emu, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

SRC, DST = sys.argv[1], sys.argv[2]

CONTENT_HEIGHT_IN = 9.69      # A4 高 11.69in - 上下页边距各 1.0in
IMAGE_BUDGET_IN = 8.30        # 一页里留给图片的总高度（其余留给 3 行标题文字）
CAP_STATEMENT = 2.40          # 题目说明图最大高度
CAP_CODE = 5.60               # 代码图最大高度
CAP_RESULT = 2.40             # 运行结果图最大高度

WP_EXTENT = qn('wp:extent')
A_EXT = qn('a:ext')

doc = Document(SRC)
body = doc.element.body
paras = doc.paragraphs

TITLE_RE = re.compile(r'^题目\s*\d+\s*[:：]\s*$')
SUB_RE = re.compile(r'^题目\s*\d+\s*(程序代码|运行结果截图)\s*[:：]\s*$')


def images_of(p_el):
    """该段落里的图片是否为 <w:drawing>（只判断，不改）"""
    return any(el.tag == WP_EXTENT for el in p_el.iter())


def size_of(p_el):
    for el in p_el.iter():
        if el.tag == WP_EXTENT:
            return int(el.get('cx')), int(el.get('cy'))
    return None


def scale_images(p_el, factor):
    for el in p_el.iter():
        if el.tag in (WP_EXTENT, A_EXT):
            cx = int(el.get('cx'))
            cy = int(el.get('cy'))
            el.set('cx', str(int(round(cx * factor))))
            el.set('cy', str(int(round(cy * factor))))


# ---------- 1) 收集段落顺序，划出题号分组 ----------
items = []
for child in list(body.iterchildren()):
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        items.append(('p', child, Paragraph(child, doc)))
    elif tag == 'tbl':
        items.append(('t', child, Table(child, doc)))

groups = {}          # 题号 -> [ (el, para) ... ]
current = None
for kind, el, obj in items:
    if kind != 'p':
        continue
    text = obj.text.strip()
    if TITLE_RE.match(text):
        current = text
        groups[current] = []
    if current is not None:
        groups[current].append((el, obj))

report = []
for title, members in groups.items():
    # ---------- 2) 删掉组内空段落（无文字、无图片） ----------
    for el, p in list(members):
        if not p.text.strip() and not images_of(el):
            el.getparent().remove(el)
            members.remove((el, p))

    # ---------- 3) 标题格式 ----------
    for el, p in members:
        t = p.text.strip()
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.0
        if TITLE_RE.match(t):
            pf.page_break_before = True
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(13)
        elif SUB_RE.match(t):
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(11)
        elif images_of(el):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------- 4) 图片等比缩放，塞进一页 ----------
    imgs = [(el, p) for el, p in members if images_of(el)]
    sizes = [size_of(el) for el, _ in imgs]
    heights = [Emu(cy).inches for _, cy in sizes]
    caps = [CAP_STATEMENT, CAP_CODE, CAP_RESULT]
    capped = [min(h, c) for h, c in zip(heights, caps)]
    total = sum(capped)
    extra = IMAGE_BUDGET_IN / total if total > IMAGE_BUDGET_IN else 1.0

    detail = []
    for (el, p), h, c in zip(imgs, heights, capped):
        factor = (c / h) * extra
        scale_images(el, factor)
        detail.append(f"{h:.2f}in -> {h*factor:.2f}in (x{factor:.3f})")
    report.append(f"{title}  图片 {len(imgs)} 张: " + " | ".join(detail) +
                  f"   合计 {sum(h*((min(h, c)/h)*extra) for h, c in zip(heights, capped)):.2f}in")

# ---------- 5) 加一行标题（学号 姓名 任务名） ----------
from docx.oxml import OxmlElement
title_el = OxmlElement('w:p')
body[0].addprevious(title_el)
title_p = Paragraph(title_el, doc)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run('example-id  Student　　第7周周三晚自习任务')
# 复制原文档第一个 run 的字体设置，保证中文字体一致
src_run = paras[0].runs[0] if paras[0].runs else None
if src_run is not None and src_run._element.rPr is not None:
    run._element.insert(0, copy.deepcopy(src_run._element.rPr))
run.bold = True
run.font.size = Pt(14)
title_p.paragraph_format.space_after = Pt(10)

doc.save(DST)
report.insert(0, f"源: {SRC}\n输出: {DST}\n可用高度 {CONTENT_HEIGHT_IN}in / 图片预算 {IMAGE_BUDGET_IN}in")
print('\n'.join(report))
