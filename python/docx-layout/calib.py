# -*- coding: utf-8 -*-
"""标定：一张 600x300px、边框贴边的图片，按 3.00x1.50in 插入，看渲染出来多大。"""
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches

im = Image.new('RGB', (600, 300), 'white')
d = ImageDraw.Draw(im)
d.rectangle([0, 0, 599, 299], outline='black', width=8)
im.save('calib.png', dpi=(96, 96))

doc = Document()
p = doc.add_paragraph()
p.add_run().add_picture('calib.png', width=Inches(3), height=Inches(1.5))
doc.add_paragraph('label 12pt').runs[0].font.size = None
doc.save('calib.docx')
print('calib.docx written')
