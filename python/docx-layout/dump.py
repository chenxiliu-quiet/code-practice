import sys
from docx import Document
from docx.shared import Emu

doc = Document(sys.argv[1])

print("=== SECTIONS ===")
for i, s in enumerate(doc.sections):
    print(f"[{i}] page {s.page_width}x{s.page_height} EMU "
          f"({Emu(s.page_width).inches:.2f}x{Emu(s.page_height).inches:.2f} in), "
          f"margins L{s.left_margin} R{s.right_margin} T{s.top_margin} B{s.bottom_margin}")

print("\n=== BODY PARAGRAPHS ===")
body = doc.element.body
from docx.table import Table
from docx.text.paragraph import Paragraph

idx_p = 0
idx_t = 0
for child in body.iterchildren():
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = Paragraph(child, doc)
        text = p.text
        has_img = 'graphic' in child.xml or 'pict' in child.xml or 'drawing' in child.xml
        marker = " [IMAGE]" if has_img else ""
        style = p.style.name
        align = p.alignment
        print(f"P{idx_p:03d} style={style!r} align={align}{marker} | {text[:180]!r}")
        idx_p += 1
    elif tag == 'tbl':
        t = Table(child, doc)
        print(f"T{idx_t:03d} TABLE rows={len(t.rows)} cols={len(t.columns)}")
        for ri, row in enumerate(t.rows):
            cells = [c.text.replace('\n', ' / ')[:60] for c in row.cells]
            print(f"      r{ri}: {cells}")
        idx_t += 1
    elif tag == 'sectPr':
        print("(sectPr)")

print(f"\ntotal paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} "
      f"inline_shapes={len(doc.inline_shapes)}")
