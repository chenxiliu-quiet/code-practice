# -*- coding: utf-8 -*-
"""一张 1000x200px（带黑色满边框）的图，按 2.00 / 3.00 / 4.00 英寸三种宽度插入。
若渲染出来三个一样宽 -> 说明渲染器按「像素数」定尺寸（忽略 wp:extent）；
若渲染出来成 2:3:4 -> 说明渲染器遵守 wp:extent。"""
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches

im = Image.new('RGB', (1000, 200), 'white')
d = ImageDraw.Draw(im)
d.rectangle([0, 0, 999, 199], outline='black', width=4)
im.save('ruler.png', dpi=(96, 96))

doc = Document()
for (w, h) in [(2.00, 0.40), (3.00, 0.60), (4.00, 0.80)]:
    p = doc.add_paragraph()
    p.add_run().add_picture('ruler.png', width=Inches(w), height=Inches(h))
    doc.add_paragraph(f'size {w} x {h} in')
doc.save('ruler.docx')
print('ruler.docx written')
